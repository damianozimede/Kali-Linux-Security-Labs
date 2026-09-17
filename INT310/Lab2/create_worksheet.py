from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

title = doc.add_heading('INT310 Lab 2: Vulnerability Assessment and Remediation', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

name = doc.add_paragraph('Patrick Damian Ozimede')
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.runs[0].bold = True

doc.add_page_break()

def add_entry(doc, cwe, filename, line, description, patch_code):
    doc.add_heading(f'{cwe} — {filename} (Line {line})', level=2)
    p = doc.add_paragraph()
    p.add_run('Description: ').bold = True
    p.add_run(description)
    p2 = doc.add_paragraph()
    p2.add_run('Patch:').bold = True
    code_para = doc.add_paragraph(patch_code)
    for run in code_para.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    doc.add_paragraph()

doc.add_heading('vulnerable_script.php', level=1)

php_entries = [
    ("CWE-89", "6-7", "The application takes the username directly from the URL ($_GET) and concatenates it into a SQL query with no sanitization, allowing an attacker to inject SQL syntax and manipulate the query.",
     "$username = $_GET['username'] ?? '';\n$stmt = mysqli_prepare($conn, \"SELECT * FROM users WHERE username = ?\");\nmysqli_stmt_bind_param($stmt, \"s\", $username);\nmysqli_stmt_execute($stmt);\n$result = mysqli_stmt_get_result($stmt);"),
    ("CWE-78", "10-11", "User-supplied $_GET['cmd'] is passed directly to system() with no validation, allowing an attacker to execute arbitrary operating system commands on the server.",
     "$cmd = $_GET['cmd'] ?? '';\n$allowed_commands = ['status'];\nif (in_array($cmd, $allowed_commands)) {\n    echo \"Command status requested.\";\n} else {\n    echo \"Command not allowed.\";\n}"),
    ("CWE-79", "17", "User input from $_POST['name'] is echoed directly onto the page with no output escaping, allowing an attacker to inject a script that executes in the victim's browser.",
     "$name = $_POST['name'] ?? '';\necho \"Hello, \" . htmlspecialchars($name, ENT_QUOTES | ENT_SUBSTITUTE, 'UTF-8');"),
    ("CWE-22 / CWE-829", "21-22", "A user-supplied path is concatenated into an include() call, allowing path traversal or inclusion of untrusted/attacker-controlled code.",
     "$path = $_GET['path'] ?? '';\n$allowed_pages = [\n    'home' => __DIR__ . '/pages/home/file.php',\n    'help' => __DIR__ . '/pages/help/file.php'\n];\nif (isset($allowed_pages[$path])) {\n    include $allowed_pages[$path];\n} else {\n    echo \"Invalid page.\";\n}"),
    ("CWE-807 / CWE-306", "26-29", "Admin access is granted purely based on a GET parameter (is_admin=true) supplied by the visitor, with no real authentication or session check performed.",
     "session_start();\n$is_admin = !empty($_SESSION['is_admin']);\nif ($is_admin) {\n    echo \"Access granted.\";\n} else {\n    echo \"Access denied.\";\n}"),
    ("CWE-798", "14-15", "The admin username and password are hardcoded directly into the source code, so anyone who can view the source can read the real credentials.",
     "$admin_user = getenv('ADMIN_USER') ?: '';\n$admin_password = getenv('ADMIN_PASSWORD') ?: '';"),
    ("CWE-434", "36-37", "An uploaded file's name and content are accepted and moved into the uploads folder with no validation on file type or extension, allowing upload of malicious code.",
     "$allowed_ext = ['jpg', 'jpeg', 'png', 'pdf'];\n$original_name = $_FILES['file']['name'];\n$ext = strtolower(pathinfo($original_name, PATHINFO_EXTENSION));\nif (in_array($ext, $allowed_ext)) {\n    $safe_name = bin2hex(random_bytes(8)) . '.' . $ext;\n    move_uploaded_file($_FILES['file']['tmp_name'], \"uploads/\" . $safe_name);\n} else {\n    echo \"File type not allowed.\";\n}"),
    ("CWE-89 / CWE-862", "41-43", "The user id used in the DELETE query comes directly from $_GET with no sanitization (SQL injection), and there is no authorization check before deleting a user account.",
     "if (isset($_GET['delete_user'])) {\n    session_start();\n    if (empty($_SESSION['is_admin'])) {\n        http_response_code(403);\n        exit('Access denied.');\n    }\n    $user_id = filter_input(INPUT_GET, 'delete_user', FILTER_VALIDATE_INT);\n    if ($user_id === false || $user_id === null || $user_id < 1) {\n        http_response_code(400);\n        exit('Invalid user ID.');\n    }\n    $stmt = mysqli_prepare($conn, \"DELETE FROM users WHERE id = ?\");\n    mysqli_stmt_bind_param($stmt, \"i\", $user_id);\n    mysqli_stmt_execute($stmt);\n}"),
    ("CWE-89 / CWE-306", "48-50", "The new password value from $_POST is concatenated directly into an UPDATE query with no sanitization, and the password change can be triggered with no authentication check.",
     "if (isset($_POST['change_password'])) {\n    session_start();\n    if (empty($_SESSION['user_id'])) {\n        http_response_code(401);\n        exit('Authentication required.');\n    }\n    $new_password = $_POST['new_password'] ?? '';\n    if (strlen($new_password) < 12 || strlen($new_password) > 128) {\n        http_response_code(400);\n        exit('Password must be 12-128 characters.');\n    }\n    $password_hash = password_hash($new_password, PASSWORD_DEFAULT);\n    $user_id = (int) $_SESSION['user_id'];\n    $stmt = mysqli_prepare($conn, \"UPDATE users SET password = ? WHERE id = ?\");\n    mysqli_stmt_bind_param($stmt, \"si\", $password_hash, $user_id);\n    mysqli_stmt_execute($stmt);\n}"),
    ("CWE-327 / CWE-759", "54", "The admin password is hashed using MD5 with no salt, a fast and cryptographically broken algorithm unsuitable for password storage.",
     "$hash = password_hash($admin_password, PASSWORD_DEFAULT);"),
    ("CWE-732", "57", "File permissions for sensitive_data.txt are set to 0777, granting read, write, and execute access to everyone on the system.",
     "chmod(\"sensitive_data.txt\", 0600);"),
    ("CWE-307", "60-62", "There is no mechanism to limit or lock out repeated failed authentication attempts.",
     "session_start();\nif (!isset($_SESSION['login_attempts'])) {\n    $_SESSION['login_attempts'] = 0;\n}\n$max_attempts = 5;\nif ($_SESSION['login_attempts'] >= $max_attempts) {\n    http_response_code(429);\n    exit(\"Too many failed login attempts. Please try again later.\");\n}\nif ($login_failed) {\n    $_SESSION['login_attempts']++;\n    echo \"Invalid credentials. Attempt \" . $_SESSION['login_attempts'] . \" of \" . $max_attempts . \".\";\n} else {\n    $_SESSION['login_attempts'] = 0;\n    echo \"Login successful.\";\n}"),
    ("CWE-134 / CWE-79", "66-67", "User input from $_GET['input'] is passed directly as the format string to printf() (format string injection), and the same unescaped output enables XSS.",
     "$input = $_GET['input'] ?? '';\n$safe_input = htmlspecialchars($input, ENT_QUOTES | ENT_SUBSTITUTE, 'UTF-8');\nprintf('%s', $safe_input);"),
    ("CWE-601", "70", "The redirect target is taken directly from $_GET['redirect_to'] with no validation, allowing redirection to an untrusted site.",
     "$redirect_to = $_GET['redirect_to'] ?? '';\n$allowed_redirects = [\n    '/' => '/',\n    '/dashboard' => '/dashboard'\n];\nif (isset($allowed_redirects[$redirect_to])) {\n    header(\"Location: \" . $allowed_redirects[$redirect_to]);\n    exit;\n} else {\n    http_response_code(400);\n    exit(\"Invalid redirect destination.\");\n}"),
    ("CWE-190", "73-74", "A value taken directly from $_GET is added to an already near-limit large integer with no bounds checking.",
     "$large_num = 999999999;\n$input_num = filter_input(INPUT_GET, 'input_num', FILTER_VALIDATE_INT);\nif ($input_num === false || $input_num === null) {\n    http_response_code(400);\n    exit(\"Invalid number.\");\n}\nif ($input_num < -1000000 || $input_num > 1000000) {\n    http_response_code(400);\n    exit(\"Number is outside the allowed range.\");\n}\n$sum = $large_num + $input_num;"),
    ("CWE-327 / CWE-759", "77", "The admin password is hashed using SHA-1 with no salt, cryptographically broken and too fast, making brute-force attacks trivial.",
     "$hash = password_hash($admin_password, PASSWORD_DEFAULT);"),
]

for cwe, line, desc, patch in php_entries:
    add_entry(doc, cwe, 'vulnerable_script.php', line, desc, patch)

doc.add_page_break()
doc.add_heading('vulnerable_script.py', level=1)

py_entries = [
    ("CWE-78", "4-5", "The program prompts for a command string and passes it directly to os.system() with no validation, allowing arbitrary OS command execution.",
     'user_input = input("Enter a command: ").strip().lower()\nallowed_commands = {\n    "status": "status"\n}\nif user_input in allowed_commands:\n    print("Command status requested.")\nelse:\n    print("Command not allowed.")'),
    ("CWE-311", "7-9", "A secret string is written directly to sensitive_data.txt in plain text with no encryption.",
     'from cryptography.fernet import Fernet\nimport os\n\nkey = os.environ.get(\'ENCRYPTION_KEY\', Fernet.generate_key())\ncipher = Fernet(key)\n\nencrypted_data = cipher.encrypt(\'Secret Data\'.encode())\nwith open(\'sensitive_data.txt\', \'wb\') as file:\n    file.write(encrypted_data)'),
    ("CWE-190", "12-17", "User-supplied numeric input is multiplied by an extremely large constant with no validation or bounds checking.",
     'MAX_NUMBER = 1_000_000\nMAX_RESULT = 10**18\n\ndef calculate(data):\n    multiplier = 1_000_000\n    result = data * multiplier\n    if abs(result) > MAX_RESULT:\n        raise ValueError("Result is too large.")\n    print(result)\n\ntry:\n    data = int(input("Enter a number: "))\n    if abs(data) > MAX_NUMBER:\n        raise ValueError("Number is outside the allowed range.")\n    calculate(data)\nexcept ValueError as e:\n    print(f"Invalid input: {e}")'),
    ("CWE-120", "20-22", "A fixed-size 10-byte buffer is created, but user input of any length is written into it via slice assignment with no length check.",
     'buffer = bytearray(10)\ndata = input("Enter more than 10 characters: ")\nencoded_data = data.encode("utf-8")\nif len(encoded_data) > len(buffer):\n    print("Input exceeds the 10-byte buffer limit.")\nelse:\n    buffer[:len(encoded_data)] = encoded_data'),
    ("CWE-807 / CWE-306", "25-29", "A privileged admin_action() function is executed purely because the user typed the word 'admin', with no real authentication.",
     'def admin_action():\n    print("Performing an admin action")\n\naction = input("Action (admin or user): ").strip().lower()\nis_authenticated_admin = False\n\nif action == "admin" and is_authenticated_admin:\n    admin_action()\nelif action == "admin":\n    print("Access denied.")'),
    ("CWE-798", "32-33", "An admin username and password are hardcoded directly into the source code as plain text.",
     'import os\n\nadmin_username = os.environ.get("ADMIN_USERNAME", "")\nadmin_password = os.environ.get("ADMIN_PASSWORD", "")'),
    ("CWE-676", "36-37", "User-supplied input is passed directly to exec(), allowing arbitrary Python code execution with the full privileges of the process.",
     'def admin_action():\n    print("Performing an admin action")\n\nexternal_func = input("Enter the function to call: ").strip()\nallowed_functions = {\n    "admin_action": admin_action\n}\nif external_func in allowed_functions:\n    allowed_functions[external_func]()\nelse:\n    print("Function not allowed.")'),
    ("CWE-79", "40-41", "User input is embedded directly into an HTML header string with no sanitization, enabling script injection if this output is ever rendered in a browser.",
     'import html\n\nuser_input = input("Enter your name: ")\nsafe_input = html.escape(user_input)\nprint(f"<h1>Hello, {safe_input}</h1>")'),
    ("CWE-494", "44-46", "A file is downloaded from a hardcoded, unencrypted (HTTP) URL and immediately executed with exec(), with no integrity check.",
     '# Remote code execution disabled: no trusted checksum source available.\nprint("Remote code execution disabled: no integrity verification available.")'),
    ("CWE-676", "49", "User-supplied input is passed directly to eval(), allowing arbitrary Python expression execution with full process privileges.",
     'import ast\n\nuser_input = input("Enter Python code to execute: ")\ntry:\n    result = ast.literal_eval(user_input)\n    print(result)\nexcept (ValueError, SyntaxError):\n    print("Invalid or disallowed input.")'),
]

for cwe, line, desc, patch in py_entries:
    add_entry(doc, cwe, 'vulnerable_script.py', line, desc, patch)

doc.save('INT310_Lab2_Worksheet_and_Patches.docx')
print('Worksheet generated')
